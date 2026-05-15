from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"D:\Code\Python")
SOURCE_MD = ROOT / "THESIS_FULL_DRAFT.md"
DEFAULT_OUTPUT_DOCX = ROOT / "中国矿业大学毕业论文_基于深度学习的同步数据捕获方法研究.docx"


TITLE = "基于深度学习的同步数据捕获方法研究"
SCHOOL = "中国矿业大学"
DOC_TYPE = "毕业设计（论文）"


def set_run_font(run, east_asia: str = "宋体", ascii_font: str = "Times New Roman", size: float | None = None, bold: bool = False):
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    run.font.name = ascii_font
    rfonts = run._element.rPr.rFonts
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(0.74)

    for style_name, size, bold in [
        ("Title", 22, True),
        ("Heading 1", 16, True),   # 三号
        ("Heading 2", 14, True),   # 四号
        ("Heading 3", 12, True),   # 小四
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体" if "Heading" in style_name else "宋体")
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.5


def clear_paragraph(paragraph) -> None:
    p = paragraph._element
    for child in list(p):
        p.remove(child)


def add_page_number(section) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    para = footer.paragraphs[0]
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    fld.append(run)
    para._p.append(fld)


def set_page_number_format(section, start: int | None = None, fmt: str | None = None) -> None:
    sectPr = section._sectPr
    existing = sectPr.find(qn("w:pgNumType"))
    if existing is not None:
        sectPr.remove(existing)
    pg = OxmlElement("w:pgNumType")
    if start is not None:
        pg.set(qn("w:start"), str(start))
    if fmt is not None:
        pg.set(qn("w:fmt"), fmt)
    sectPr.append(pg)


def add_manual_toc(doc: Document, entries: list[tuple[int, str]]) -> None:
    for level, text in entries:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0 if level == 1 else 0.9)
        run = p.add_run(text)
        set_run_font(run, size=12 if level == 1 else 11, bold=(level == 1))


def add_cover_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    run = p.add_run(SCHOOL)
    set_run_font(run, east_asia="黑体", size=24, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DOC_TYPE)
    set_run_font(run, east_asia="黑体", size=22, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(TITLE)
    set_run_font(run, east_asia="黑体", size=20, bold=True)

    info_items = [
        "学生姓名：________________",
        "学    号：________________",
        "学    院：________________",
        "专    业：________________",
        "指导教师：________________",
        "完成时间：2026年5月",
    ]
    for item in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(4.8)
        if item == info_items[0]:
            p.paragraph_format.space_before = Pt(70)
        run = p.add_run(item)
        set_run_font(run, size=14)

    doc.add_page_break()


def strip_md_links(text: str) -> str:
    return re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)


def html_like_inline(text: str) -> str:
    text = strip_md_links(text)
    text = text.replace("`", "")
    return text.strip()


def add_paragraph(doc: Document, text: str, style: str | None = None, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent: bool = True) -> None:
    p = doc.add_paragraph(style=style)
    p.alignment = align
    if style is None and first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    else:
        p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(html_like_inline(text))
    if style == "Heading 1":
        set_run_font(run, east_asia="黑体", size=16, bold=True)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
    elif style == "Heading 2":
        set_run_font(run, east_asia="黑体", size=14, bold=True)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
    elif style == "Heading 3":
        set_run_font(run, east_asia="黑体", size=12, bold=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
    elif style == "Title":
        set_run_font(run, east_asia="黑体", size=18, bold=True)
    else:
        set_run_font(run, size=12)


def add_bullet(doc: Document, text: str) -> None:
    cleaned = html_like_inline(text)
    if not cleaned:
        return
    p = doc.add_paragraph(style=None)
    p.style = doc.styles["Normal"]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.hanging_indent = Cm(0.74)
    run = p.add_run("• " + cleaned)
    set_run_font(run, size=12)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(html_like_inline(text))
    set_run_font(run, size=10.5, bold=bold)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            set_cell_text(table.cell(r_idx, c_idx), text, bold=(r_idx == 0))


def parse_table(lines: list[str]) -> list[list[str]]:
    parsed: list[list[str]] = []
    for idx, line in enumerate(lines):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if idx == 1 and all(set(c) <= {"-", ":"} for c in cells):
            continue
        parsed.append(cells)
    return parsed


def add_image(doc: Document, alt: str, path_str: str) -> None:
    normalized = re.sub(r"^/([A-Za-z]:/)", r"\1", path_str)
    path = Path(normalized)
    if not path.exists():
        add_paragraph(doc, f"{alt}（图片文件缺失：{path}）", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.5))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.first_line_indent = Cm(0)
    run = cap.add_run(alt)
    set_run_font(run, size=10.5)


def add_running_header(section, text: str) -> None:
    section.header.is_linked_to_previous = False
    para = section.header.paragraphs[0]
    clear_paragraph(para)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run_font(run, size=10.5)


def build_doc(output_docx: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_cover_page(doc)

    cover_section = doc.sections[0]
    cover_section.different_first_page_header_footer = True
    cover_section.footer.is_linked_to_previous = False
    clear_paragraph(cover_section.footer.paragraphs[0])

    front_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_document(doc)
    front_section.footer.is_linked_to_previous = False
    clear_paragraph(front_section.footer.paragraphs[0])

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    toc_entries: list[tuple[int, str]] = []
    for raw in lines:
        s = raw.rstrip()
        if s.startswith("# ") and s[2:].strip() not in {"摘要", "Abstract", "目录", "参考文献", "致谢"}:
            toc_entries.append((1, s[2:].strip()))
        elif s.startswith("## "):
            toc_entries.append((2, s[3:].strip()))
    paragraph_buffer: list[str] = []
    table_buffer: list[str] = []
    pending_figure_caption: str | None = None
    in_toc = False
    current_top_heading = ""

    def flush_paragraph():
        nonlocal paragraph_buffer, pending_figure_caption, current_top_heading
        if paragraph_buffer:
            text = "".join(s.strip() for s in paragraph_buffer).strip()
            if text:
                if re.match(r"^图\d", text):
                    pending_figure_caption = text
                else:
                    pending_figure_caption = None
                    if current_top_heading == "Abstract":
                        add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
                    else:
                        add_paragraph(doc, text)
            paragraph_buffer = []

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, parse_table(table_buffer))
            table_buffer = []

    first_h1 = True
    image_re = re.compile(r"!\[(.*?)\]\((.*?)\)")

    for raw in lines:
        line = raw.rstrip().lstrip("\ufeff")

        if line.startswith("|"):
            flush_paragraph()
            table_buffer.append(line)
            continue
        flush_table()

        if not line.strip():
            flush_paragraph()
            continue

        m = image_re.match(line.strip())
        if m:
            flush_paragraph()
            caption = pending_figure_caption or m.group(1)
            pending_figure_caption = None
            add_image(doc, caption, m.group(2))
            continue

        if line.startswith("# "):
            flush_paragraph()
            title = line[2:].strip()
            if title == "目录":
                in_toc = True
                current_top_heading = title
                add_paragraph(doc, title, style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False)
                add_manual_toc(doc, toc_entries)
                doc.add_page_break()
                continue
            if title in {"Abstract"} and not first_h1:
                doc.add_page_break()
            if in_toc:
                in_toc = False
                main_section = doc.add_section(WD_SECTION.NEW_PAGE)
                add_page_number(main_section)
                set_page_number_format(main_section, start=1, fmt="decimal")
                add_running_header(main_section, TITLE)
            elif not first_h1:
                if title in {"摘要", "Abstract", "目录"}:
                    doc.add_page_break()
                elif not re.match(r"^第\d+章", title):
                    doc.add_page_break()
            first_h1 = False
            current_top_heading = title
            align = WD_ALIGN_PARAGRAPH.CENTER if title in {"摘要", "Abstract", "目录", "参考文献", "致谢"} or re.match(r"^第\d+章", title) else WD_ALIGN_PARAGRAPH.LEFT
            add_paragraph(doc, title, style="Heading 1", align=align, first_indent=False)
            continue
        if in_toc:
            continue
        if line.startswith("## "):
            flush_paragraph()
            add_paragraph(doc, line[3:].strip(), style="Heading 2", align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
            continue
        if line.startswith("### "):
            flush_paragraph()
            add_paragraph(doc, line[4:].strip(), style="Heading 3", align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False)
            continue
        if line.startswith("- "):
            flush_paragraph()
            add_bullet(doc, line[2:].strip())
            continue

        if line.startswith("关键词：") or line.startswith("Key Words:"):
            flush_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(0)
            run = p.add_run(line.strip())
            set_run_font(run, size=12, bold=True)
            continue

        paragraph_buffer.append(line)

    flush_paragraph()
    flush_table()
    doc.save(str(output_docx))


if __name__ == "__main__":
    output = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT_DOCX
    build_doc(output)
    print(output)
